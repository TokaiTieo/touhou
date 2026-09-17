# -*- coding: utf-8 -*-
"""将本次对话整理为 Word 文档。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# 基础字体
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return p

def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")

def qa(q, a_points):
    p = doc.add_paragraph()
    r = p.add_run("问:" + q)
    r.bold = True
    for point in a_points:
        bullet(point)

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = htext
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return t

# ================= 封面 =================
title = doc.add_heading("《东方异变录》项目工作与技术讨论记录", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("日期:2026-08-19    版本:v0.12.0(存档结构 V7)")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

# ================= 第一部分 =================
h("第一部分 前端唯美主题视觉重做(已交付)", 1)
para("本次将前端整体视觉从「墨黑暗色」迁移为「薄樱·和纸·金繕·月青·藤紫」的东方 Project 唯美浅色主题。"
     "实现方式为纯 CSS 追加覆写:仅在 css/app.css 与 css/vue-game.css 两文件末尾新增覆写层,"
     "Vue 模板、组件拆分、API 契约、LangGraph 编排、V7 存档全部零改动。")

h("1.1 全局设计令牌", 2)
table(["类别", "变量", "值"], [
    ["和纸", "--washi / --washi-warm", "#FBF8F1 / #F5EDE2"],
    ["淡墨", "--ink / --ink-soft / --ink-faint", "#2A2A2A / #574D54 / #9A8D93"],
    ["樱", "--sakura / --sakura-deep", "#F8C4D5 / #DF89A5"],
    ["金", "--kin / --kin-bright", "#B78F45 / #E6CD94"],
    ["月青", "--tsuki / --tsuki-deep", "#A8D8E8 / #5E96AC"],
    ["藤紫", "--fuji / --fuji-deep", "#B19CD8 / #8A72B4"],
    ["柔朱", "--shu", "#C05A68"],
    ["圆角", "--radius-s/m/l", "8px / 12px / 16px"],
    ["动效", "--ease-drift", "cubic-bezier(.25,.6,.3,1),300ms+"],
    ["字体", "--font-mincho", "Yu Mincho→MS Mincho→SimSun→Songti SC"],
])

h("1.2 五个核心模块改造要点", 2)
bullet("对话流:和纸白气泡+金色左边线;玩家樱粉渐变、NPC 月青、系统藤紫胶囊;圆形樱印;行距 1.9。")
bullet("地图:地点改为「境界札」圆角白卡,当前地点樱粉边+樱雾渐变;危险度细线胶囊三档。")
bullet("NPC 与缘分录:圆形头像缘分色描边;关系值藤紫;保持列表式(图谱可视化未做)。")
bullet("状态/任务:异变进度条改胶囊槽+樱粉/金渐变;任务金札钉;工具栏徽章胶囊化。")
bullet("设置/启动:启动画面「月下樱散」缓漂樱瓣;英雄横幅晨光渐变+月轮光晕;弹窗 16px 卷轴。")

h("1.3 第二幕打磨", 2)
bullet("全部红色实底印章改为白底朱文细描边印(含启动屏「東/方」、界/記/遇/人等)。")
bullet("对话舞台/侧栏/气泡/顶栏引入 backdrop-filter 磨砂玻璃;细轨圆润滑动条。")
bullet("修复「宣言行动」按钮折行;默认主题薄樱、light-theme 晴日双浅色变体。")

h("1.4 验证与产物", 2)
bullet("Playwright e2e 7/7 通过(含可玩闭环、无副作用改写、V7 分支重载、桌面/移动视觉截图)。")
bullet("Python 测试 107 通过;5 项关系服务测试为改动前既有失败,与 CSS 无关。")
bullet("重新打包 touhou.exe(230,619,421 bytes)并替换根目录;冒烟测试通过。")
bullet("重新生成分发测试包 release/touhou-test-package.zip(白名单,空 Key,无存档)。")
bullet("已提交并推送 GitHub main:4439f90 feat(ui): 唯美主题视觉重做。")

# ================= 第二部分 =================
doc.add_page_break()
h("第二部分 项目工程问答(架构机制)", 1)

h("2.1 超时 / 心跳 / 判活", 2)
qa("项目有超时/心跳判活机制吗?hang 会无限阻塞吗?", [
    "已有多层有限超时:AI HTTP 60s(最多4次重试+跨模型 fallback)、流式 60s 无 chunk 超时、asyncio wait_for 65s、LangGraph task 75s/entrypoint 95s、前端断线轮询 90s 硬上限。",
    "模型慢/断流/断连/进程崩溃四类常见故障均能有限等待后报错或自动恢复,不会无限阻塞。",
    "唯一无限阻塞场景:后端进程/事件循环整体假死——uvicorn 无请求超时、桌面端启动后无看门狗、前端普通 fetch 无超时。",
    "次要风险:SSE 无心跳;wait_for 超时后 executor 线程不可回收;锁获取无超时(目前靠 AI 层超时兜底)。",
    "完整风险清单已存档至 docs/备忘录-超时与判活风险.txt。",
])

h("2.2 记忆召回(n-gram)", 2)
qa("n-gram 召回几条消息?", [
    "NPC 对话:该 NPC 最多 8 条记忆,有长期印象则前置一行摘要。",
    "环境行动:每个 NPC 各召回最多 2 条,最终保留最后 8 个 NPC 的记忆段。",
    "召回后还要过 prompt 分区预算(npc_memories 3000 字符,超预算保尾截断)。",
])
qa("n-gram 的权重是 0.3/0.3/0.3 吗?为什么不用等权?", [
    "实际为 1-gram 0.25、2-gram 1.0、3-gram 0.7(memory_retrieval.py:103),另有概念组命中加 2.5+命中数,向量做 L2 归一化。",
    "中文双字词占比最高、判别力最强,故 2-gram 拉满;单字噪音大压到 0.25;3-gram 稀疏给 0.7。",
    "等权会让判别力最弱的单字特征在归一化向量中嗓门最大,稀释真正的语义匹配。",
])
qa("召回用的是 RRF 吗?", [
    "不是 RRF,是加权线性融合 + MMR 去重:语义×24+稠密×18+关键词×3.5+重要度×1.6+使用次数+时间衰减+置信度+真伪调整(superseded -8),选取时与已选条目相似度×5 惩罚。",
    "RRF 适合多独立检索器;本项目只有 1-2 个检索信号加业务元数据,线性融合更直接可控。",
])

h("2.3 Prompt 拼接与上下文布局", 2)
qa("召回的记忆在 prompt 里怎么放?", [
    "单独开一段(## 【当前NPC记忆】),不与历史混排、不按相关度穿插。",
    "模板顺序:世界观→NPC/玩家信息→场景→对话历史(近20轮)→长期剧情摘要→NPC态度→NPC记忆→物品/线索→玩家输入→规则。",
    "记忆段位于历史之后、玩家输入之前的近因优势区;段内按召回得分排序,长期印象置顶。",
])
qa("为什么先放最近 20 轮再放摘要?会不会注意力漂移?", [
    "LLM 长上下文注意力呈 U 形,中段是洼地;模板把静态身份锚点放开头、决策相关状态与玩家输入放尾部,历史(长文本)主动承担中段洼地。",
    "摘要段开头有仲裁指令「若与最近对话冲突,以最近对话为准」,该指令要求历史先被阅读,故历史在前、摘要在后。",
    "真实风险:窗口内中间轮次半遗忘、第 21 轮刚挤出窗口而摘要未重建的真空期(靠关键词即时触发摘要缓解)。",
])
qa("prompt 如何拼接?", [
    "四步:各分区文本生成→budget_context_sections 预算分配(总 24000 字符,插入顺序即优先级,protected 三区优先)→render_prompt 用 str.replace 替换具名占位符(刻意不用 .format,模板含字面 JSON 契约示例)→尾部追加规则预裁定/叙事导演/世界回响/离屏动向。",
])

h("2.4 关键事件判定与摘要压缩", 2)
qa("用什么 prompt 让 LLM 判定关键事件?每次输入都多跑一次吗?误判怎么办?", [
    "没有专门判定 prompt。剧情摘要触发是纯本地关键词(IMPORTANT_WORDS 15 词);NPC 记忆靠主响应契约的 memory_updates 字段顺带返回,每回合只有 1 次 LLM 调用。",
    "后端 build_auto_memory_updates 确定性兜底:关系变化(8)/符卡(8)/线索(6)/事件(6)/互动关键词(7),LLM 漏判不丢硬事实。",
    "误判防护:fact_key 冲突按置信度判 disputed/superseded(召回 -8 沉底)、Pydantic 契约校验丢弃非法字段、记忆只影响叙事不改数值、制作人控制台可增删改。",
])
qa("记忆分几层?每层何时写入、何时触发摘要压缩?压缩后去哪、原文还在吗?", [
    "四层:①conversation_history 全量对话(每回合追加,永不删)②story_summary 剧情摘要 ③npc_memories 记忆桶 ④npc_memory_summaries 印象摘要。",
    "②触发:历史新增≥12 条/最新消息命中关键词/历史变短或 force;从当前历史确定性重建,整体替换。",
    "③→④触发:单 NPC 桶>30 条,每次写入后检查;保留最近 24 条,旧条目按重要度 top6 熔进 900 字印象。",
    "去向:②摘要不删原文(历史全量留档);④摘要换原文(旧条目丢弃,唯一备份是 40 节点快照)。",
])
qa("摘要本身是谁做的?同步还是异步?失败会断上下文吗?", [
    "摘要是本地确定性字符串操作,不调大模型(文件头注明 without an extra AI request),同步执行但耗时微秒级。",
    "不存在超时/格式失败面;唯一可失败的是主 LLM 调用本身,此时整个回合原子失败、什么都不写入。",
    "原始对话永不因摘要删除(最近 20 轮只是注入窗口),上下文不会断。",
])
qa("独立记忆桶是什么?", [
    "npc_memories 按 NPC 分桶隔离:存储、检索、向量索引、压缩四个维度各自独立;每桶一个数组,召回按桶进行。",
    "压缩触发:单桶>30 条;算法:切(留近24)→选(旧条目重要度top6)→熔(拼进900字印象)→弃(旧条目)。",
    "小不一致:压缩选 top6 不过滤 superseded(召回侧有 -8 惩罚,压缩侧没有),可一行修复。",
])

h("2.5 会话记忆持久化与删除连续性", 2)
qa("会话记忆如何保存?", [
    "唯一事实源:worlds/world_touhou/sessions/characters/{uuid}.json(+分离的 {uuid}_tasks.json)。",
    "回合中全部变化在内存结算,回合末 save_turn_bundle 一次提交;临时文件+os.replace 原子写入;逐角色锁+事务日志;turn_id 幂等;state_revision 乐观并发。",
    "LangGraph checkpoint(runtime/turn_checkpoints.sqlite3)仅为被中断回合服务,成功即删,不属存档。",
])
qa("删除对话后如何保证存档连续性、不受已删内容影响?", [
    "快照恢复/分支=真·时间线回滚:每回合自动创建完整角色+任务联合快照(40 节点),恢复即整体覆盖,被删回合的记忆/关系/任务/异变全部消失。",
    "delete_history 只截断对话历史并 force 重建 story_summary(防止摘要泄漏被删内容),但不回滚 npc_memories/关系/任务——要彻底回滚须用快照恢复。",
    "配套:历史变短自动强制重建摘要;改写零副作用;state_revision 防旧在途回合幽灵写入;turn_id 幂等防重复结算。",
])

h("2.6 检索与模型组件", 2)
qa("项目用到 embedding 和 rerank 吗?", [
    "Embedding:默认是本地字符 n-gram 稀疏向量(非神经网络);sentence-transformers 稠密向量为预留通道(未安装未配置);sqlite_vec 仅在打包 spec 的 hiddenimports 中,后端零调用。",
    "Rerank:无独立 reranker;召回即加权线性融合排序,MMR 多样性惩罚分担 rerank 职责。候选集小(单 NPC≤30 条),两段式属过度工程。",
])
qa("后端没有 turn_id 是谁补?coordinator 还是 runner?", [
    "补全规则 ensure_turn_id 定义在 coordinator(空则生成 UUID);实际补全动作发生在 endpoint 装饰器层(TurnRunner.endpoint 现行路径、TurnCoordinator.endpoint 旧路径),在进锁与 in-flight 去重之前补齐并回写 request.turn_id。",
])

h("2.7 关系系统", 2)
qa("关系图谱是如何实现的?", [
    "非图数据库、非图可视化:以玩家为中心的星型边集合,文本态度+数值进度双轨(relationships_map / relationship_progress / relationships_history 20 条 / relationship_boundaries 冷却)。",
    "八档阶段:死敌≤-75/敌对/疏离/相识/友好/信赖/亲密/恋人≥88;LLM 返回文本态度经关键词映射入分。",
    "节奏钳制:普通每回合最多+12,明确示好+18,≥70 无关系词+8,拒绝冷却期只能降,恶化最多-30,制作人直接覆盖;钳制记录 requested_stage 与 pacing_reason。",
    "前端为「缘分录」账簿列表,无节点连线图;无 NPC-NPC 边。",
])

# ================= 第三部分 =================
doc.add_page_break()
h("第三部分 运维事项记录", 1)
h("3.1 GitHub", 2)
bullet("仓库 https://github.com/TokaiTieo/touhou 已配置远程,提交 4439f90 已推送 main。")
bullet("本机 push 免确认的原因:Git Credential Manager 已将 GitHub 令牌存入 Windows 凭据管理器(首次 OAuth 授权后),进程以本机用户身份运行自动取凭证。")
bullet("收紧方式:git credential-manager erase 清除凭证;GitHub 网页吊销 OAuth App;或添加 ruleset。")
bullet("Ruleset 已配置(名称 revision,Active):Target=Default branch;规则=禁删分支+禁 force push+需 PR(0 评审)+线性历史;Bypass=Repository admin(always)。经 API 核实生效;因 bypass 包含本人,规则实际约束未来协作者。")
h("3.2 打包产物", 2)
bullet("python -m PyInstaller --noconfirm --clean api_release.spec;touhou.exe 230,619,421 bytes,SHA-256 610098949A7259C31B065D5A04328980432F27DF971ED2AC60CCCEBBE01F9EF3。")
bullet("scripts/smoke-exe.ps1 冒烟通过;release/touhou-test-package.zip 白名单重打包(仅 EXE+脚本+玩前必读+空 Key .env)。")
bullet("sentence-transformers 为可选稠密向量后端(env: TOUHOU_EMBEDDING_MODEL / TOUHOU_ALLOW_MODEL_DOWNLOAD),默认 local_ngram,EXE 永远走 local_ngram,建议不启用。")

h("附:题外问答(米哈游日常实习笔试)", 2)
bullet("官方校招/全年实习流程含笔试(简历→笔试→面试→offer),但存在免笔试通道(提前批/线下直通车/部分岗位直进面试)。")
bullet("真正的日常实习独立于校招项目、无毕业时间限制,通常以简历+业务面试为主,无统一集中笔试;技术岗可能有现场做题或作品任务。以具体岗位 JD 与 HR 通知为准。")

doc.save("docs/东方异变录-工作与技术讨论记录-2026-08-19.docx")
print("saved")
